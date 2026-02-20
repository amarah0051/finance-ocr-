import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.models.schemas import ExtractedDocument


def _load_mapping() -> dict:
    mapping_path = Path(__file__).resolve().parents[1] / "config" / "mapping_rules.json"
    with mapping_path.open("r", encoding="utf-8") as fh:
        return json.load(fh)


def _validate_headers(ws, header_row: int, expected_headers: list[str], strict: bool) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col).value
        if value in expected_headers:
            header_map[value] = col

    if strict:
        missing = [h for h in expected_headers if h not in header_map]
        if missing:
            raise ValueError(f"Strict mapping failed. Missing Excel headers: {missing}")

    return header_map


def _find_next_row(ws, header_row: int, col_index: int) -> int:
    r = header_row + 1
    while ws.cell(r, col_index).value not in (None, ""):
        r += 1
    return r


def _eu_date_or_none(value: str | None) -> str | None:
    if not value:
        return None
    for fmt in ("%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(value, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue
    return value


def append_to_workbook(template_or_existing: Path, output_path: Path, doc: ExtractedDocument) -> int:
    mapping = _load_mapping()
    wb = load_workbook(template_or_existing)

    sheet_name = mapping["sheet_name"]
    ws = wb[sheet_name]

    header_row = int(mapping["header_row"])
    expected_headers = list(mapping["columns"].keys())
    strict = bool(mapping.get("strict", True))
    header_map = _validate_headers(ws, header_row, expected_headers, strict)

    first_col = header_map["Internal ID"]
    next_row = _find_next_row(ws, header_row, first_col)

    for item in doc.items:
        row_values = {
            "Internal ID": doc.internal_id,
            "Title": doc.title,
            "Currency": doc.currency,
            "Date": _eu_date_or_none(doc.date),
            "Reseller": doc.reseller,
            "ResellerContact": doc.reseller_contact,
            "Expires": _eu_date_or_none(doc.expires),
            "ExpectedClose": _eu_date_or_none(doc.expected_close),
            "EndUser": doc.end_user,
            "BusinessUnit": doc.business_unit,
            "Item": item.item,
            "Quantity": item.quantity,
            "Salesprice": item.list_unit_price,
            "Purchaseprice": item.net_unit_price,
            "PurchaseDiscount": item.discount_percent,
            "Location": doc.location,
            "ContractStart": _eu_date_or_none(item.contract_start),
            "ContractEnd": _eu_date_or_none(item.contract_end),
            "Serial#Supported": doc.serial_supported,
            "Rebate": doc.rebate,
            "Opportunity": doc.opportunity,
        }

        for header, value in row_values.items():
            ws.cell(next_row, header_map[header], value=value)

        # Backend formula logic for Salesdiscount for every appended row.
        formula = f"=1-(((O{next_row}/P1)*(1+P2)/M{next_row}))"
        ws.cell(next_row, header_map["Salesdiscount"], value=formula)

        next_row += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return len(doc.items)
